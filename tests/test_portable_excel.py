import importlib.util
import io
from pathlib import Path
import pytest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]

@pytest.mark.parametrize('app',['retail','onboarding','catalog','po'])
def test_export_preserves_all_rows_numbers_and_literal_text(app):
    spec=importlib.util.spec_from_file_location('export_'+app,ROOT/'apps'/app/'portable_excel.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    rows=[{'SKU':f'DEMO-{i:04d}','Quantity':60,'Cost':650.5,'Note':'=1+1','Missing':None} for i in range(600)]
    payload=module.excel_bytes({'PO / lines':rows,'Empty':[],'Context':[{'Population':600,'Definition':'Synthetic test rows'}]})
    book=load_workbook(io.BytesIO(payload),data_only=False)
    sheet=book['PO _ lines']
    assert sheet.max_row==601 and sheet['A601'].value=='DEMO-0599'
    assert sheet['B2'].value==60 and sheet['C2'].value==650.5
    assert sheet['D2'].value=='=1+1' and sheet['D2'].data_type=='s'
    assert sheet['E2'].value is None
    assert sheet.freeze_panes=='A2' and sheet.auto_filter.ref=='A1:E601'
    assert book['Context']['A2'].value==600
    assert book['Empty']['A1'].value=='No records'

@pytest.mark.parametrize('app',['retail','onboarding','catalog','po'])
def test_empty_workbook_and_colliding_names(app):
    spec=importlib.util.spec_from_file_location('empty_'+app,ROOT/'apps'/app/'portable_excel.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    assert load_workbook(io.BytesIO(module.excel_bytes({}))).sheetnames==['No records']
    names=['a'*40,'a'*39+'b','A'*40,'[]:*?/\\']
    book=load_workbook(io.BytesIO(module.excel_bytes({name:[] for name in names})))
    assert len(book.sheetnames)==4
    assert len({name.lower() for name in book.sheetnames})==4
    assert all(len(name)<=31 for name in book.sheetnames)

