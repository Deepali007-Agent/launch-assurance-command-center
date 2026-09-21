import sys,os
from pathlib import Path
p=Path(__file__).resolve().parents[1]/'streamlit_app.py'
original=os.mkdir
os.mkdir=lambda path,mode=0o777,**kw:original(path,0o777 if mode==0o700 else mode,**kw)
from streamlit.testing.v1 import AppTest
at=AppTest.from_file(str(p),default_timeout=90).run()
assert not at.exception,at.exception
[b for b in at.button if b.label=='Run launch assessment'][0].click().run()
assert not at.exception,at.exception
assert next(m.value for m in at.metric if m.label=='Ready for review')=='2 / 4 SKUs'
assert next(m.value for m in at.metric if m.label=='Potential sales recovered')=='INR 187,500'
bt=AppTest.from_file(str(p),default_timeout=90).run()
assert not bt.exception,bt.exception
assert not bt.metric
assert 'plan' not in bt.session_state['_la_demo_state']
print('Public demo PASS: source validation, financial output, independent visitor state.')

for entry in at.text_input:
 if entry.label=='Simulation reviewer':entry.input('Portfolio QA')
[b for b in at.button if b.label=='Simulate selected actions & revalidate'][0].click().run()
assert not at.exception,at.exception
assert next(m.value for m in at.metric if m.label=='Potential sales recovered')=='INR 712,500'
assert next(m.value for m in at.metric if m.label=='Potential net margin improvement')=='INR 397,950'
export_button=[b for b in at.button if b.label=='Prepare Excel'][-1]
export_key='xlsx_'+export_button.key.removeprefix('prepare_')
export_button.click().run()
assert not at.exception,at.exception
from io import BytesIO
from openpyxl import load_workbook
exports=[at.session_state[export_key]]
assert exports
book=load_workbook(BytesIO(exports[0]),data_only=True)
assert book['Evidence'].max_row==3
assert bt.session_state['_la_demo_state']['history']==[]
print('Revalidation and commercial Excel readback PASS; second visitor remains unchanged.')

