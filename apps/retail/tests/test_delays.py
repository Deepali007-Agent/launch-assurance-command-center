import pytest
from retail_workflow.store import WorkflowStore

@pytest.fixture
def setup(tmp_path):
    store=WorkflowStore(tmp_path/'workflow.db');batch=store.create_batch('Launch','QC')
    vendor={'vendor_id':'V1','legal_name':'Vendor','currency':'INR'}
    item={'sku':'S1','vendor_id':'V1','product_name':'Shirt','brand':'Brand','category':'Apparel','gender':'Men','price':100,'description':'Original content'}
    store.ingest_onboarding(batch,[vendor],[item],{'V1':{'status':'Clear','findings':[]}},'QC')
    store.complete_catalog(store.catalog_queue(batch),{'S1':{'status':'Clear','findings':[]}})
    po={'PO_ID':'P1','SKU':'S1','Vendor ID':'V1','Class':'Apparel','gender':'Men'}
    result={'status':'PASS','errors':[],'warnings':[]}
    store.ingest_pos(batch,[po],[result],'QC')
    return store,batch,vendor,item,po,result


def finish(store,batch,kind):
    row=next(r for r in store.board(batch) if r['kind']==kind)
    store.approve([(row['id'],row['etag'])],'Reviewer','Test')
    row=next(r for r in store.board(batch) if r['kind']==kind)
    store.simulate_erp([(row['id'],row['etag'])],'Operator','Test')



from retail_workflow.delays import analyze,record_attribution

def test_complete_data_can_still_wait_for_finance(setup):
    s,b,*_=setup
    ready,delays=analyze(s,b)
    assert bool(ready.iloc[0]['Item data ready'])
    assert not bool(ready.iloc[0]['Can proceed to simulated item setup'])
    finance=delays[delays['Accountable team']=='Finance'].iloc[0]
    assert finance['Originating source']=='Not established'
    assert finance['Recorded finance reason']=='Not provided'
    assert finance['Accountable owner']=='Unassigned to this team'
    assert finance['Linked PO lines']==1

def test_explicit_attribution_does_not_approve_or_change_origin_on_reassign(setup):
    s,b,*_=setup
    _,rows=analyze(s,b);row=rows[rows['Accountable team']=='Finance'].iloc[0]
    with pytest.raises(ValueError):record_attribution(s,b,row['Delay key'],'Vendor submission','Outstanding dues reported','','Recorder','')
    record_attribution(s,b,row['Delay key'],'Vendor submission','Outstanding dues reported','FIN-TICKET-123','Finance reviewer','Documented test evidence')
    vendor=s.requests(b,'vendor')[0]
    s.assign(vendor['id'],vendor['etag'],'Finance','Finance owner','2026-12-01','Open','Manager','Route to Finance')
    _,rows=analyze(s,b);updated=rows[rows['Delay key']==row['Delay key']].iloc[0]
    assert updated['Originating source']=='Vendor submission'
    assert updated['Accountable owner']=='Finance owner'
    assert not any(r['Approved'] for r in s.board(b))

def test_old_finance_attribution_not_carried_to_new_vendor_version(setup):
    s,b,v,i,*_=setup
    _,rows=analyze(s,b);row=rows[rows['Accountable team']=='Finance'].iloc[0]
    record_attribution(s,b,row['Delay key'],'Finance','Outstanding dues reported','REF','Recorder','')
    s.ingest_onboarding(b,[dict(v,legal_name='Revised')],[i],{'V1':{'status':'Clear','findings':[]}},'QC')
    _,rows=analyze(s,b)
    assert all(rows['Recorded finance reason']=='Not provided')
    with pytest.raises(ValueError,match='no longer current'):record_attribution(s,b,row['Delay key'],'Finance','Approval review pending','REF','Recorder','')

def test_gates_clear_after_appropriate_review(setup):
    s,b,*_=setup
    finish(s,b,'vendor');finish(s,b,'item')
    ready,rows=analyze(s,b)
    assert bool(ready.iloc[0]['Can proceed to simulated item setup'])
    assert rows.empty

def test_non_finance_delay_cannot_claim_financial_reason(setup):
    s,b,*_=setup
    _,rows=analyze(s,b);row=rows[rows['Accountable team']!='Finance'].iloc[0]
    with pytest.raises(ValueError,match='Finance dependency'):record_attribution(s,b,row['Delay key'],'Finance','Outstanding dues reported','REF','Recorder','')


def test_register_contains_ready_and_waiting_skus_and_team_details(setup):
    from retail_workflow.clarity import workbook_sheets,primary_step
    from retail_workflow.dashboard import evidence
    from portable_excel import excel_bytes
    from openpyxl import load_workbook
    from io import BytesIO
    s,b,*_=setup
    sheets=workbook_sheets(s,b)
    assert len(sheets['SKU overview'])==1
    assert sheets['SKU overview'][0]['SKU']=='S1'
    assert sheets['Finance'][0]['SKU']=='S1'
    assert sheets['Finance'][0]['Next action owner']=='Unassigned to this team'
    row=next(r for r in s.board(b) if r['kind']=='item')
    assert primary_step(row,evidence(s,b))=='Awaiting Finance approval'
    book=load_workbook(BytesIO(excel_bytes(sheets)))
    assert book['SKU overview'].max_row==2
    assert book['Requirements'].max_row==len(sheets['Requirements'])+1
    finish(s,b,'vendor');finish(s,b,'item')
    updated=workbook_sheets(s,b)
    assert len(updated['SKU overview'])==1
    assert updated['Requirements']==[]
    assert updated['SKU overview'][0]['Primary next step']=='Setup acknowledged (simulated)'


def test_missing_catalog_validation_is_named_not_generic_blocked(setup):
    from retail_workflow.clarity import primary_step
    from retail_workflow.dashboard import evidence
    s,b,v,i,*_=setup
    s.ingest_onboarding(b,[v],[dict(i,product_name='Revised')],{'V1':{'status':'Clear','findings':[]}},'QC')
    row=next(r for r in s.board(b) if r['kind']=='item')
    assert primary_step(row,evidence(s,b))=='Awaiting Catalog validation'


def test_actionable_reviews_separate_from_waiting_setup(setup):
    s,b,v,i,*_=setup
    _,rows=analyze(s,b)
    finance=rows[rows['Accountable team']=='Finance'].iloc[0]
    assert finance['Action status']=='Action needed now'
    approval=rows[rows['Requirement']=='Current version needs a recorded human approval'].iloc[0]
    assert approval['Action status']=='Action needed now'  # Catalog review can run in parallel.
    setup_row=rows[rows['Requirement'].str.contains('Vendor master')].iloc[0]
    assert setup_row['Action status']=='Waiting on another step'
    s.ingest_onboarding(b,[v],[i,dict(i,sku='S2')],{'V1':{'status':'Clear','findings':[]}},'QC')
    _,rows=analyze(s,b)
    finance=rows[rows['Accountable team']=='Finance']
    assert finance.SKU.nunique()==2 and finance['Task ID'].nunique()==1
    assert set(finance['Actionable elapsed time'])=={'Not measured'}
    s.ingest_onboarding(b,[v],[i],{'V1':{'status':'Blocked','findings':['Missing tax ID']}},'QC')
    _,rows=analyze(s,b)
    assert set(rows[rows['Accountable team']=='Finance']['Action status'])=={'Waiting on another step'}
