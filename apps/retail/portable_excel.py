"""In-memory Excel export using the declared openpyxl dependency."""
import io
import json
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def excel_bytes(sheets):
    book=Workbook();book.remove(book.active)
    used=set()
    for label,records in sheets.items():
        title=re.sub(r'[\\/*?:\[\]]','_',str(label)).strip("'")[:31] or 'Data'
        base=title;number=1
        while title.casefold() in used:
            number+=1;tail=' '+str(number);title=base[:31-len(tail)]+tail
        used.add(title.casefold());sheet=book.create_sheet(title)
        records=list(records)
        columns=list(dict.fromkeys(key for row in records for key in row))
        sheet.append(columns or ['No records'])
        for row_number,row in enumerate(records,start=2):
            sheet.append([json.dumps(row.get(key),default=str) if isinstance(row.get(key),(dict,list,tuple)) else row.get(key) for key in columns])
            # Uploaded text must remain literal text, never executable formulas.
            for column_number in range(1,len(columns)+1):
                cell=sheet.cell(row_number,column_number)
                if isinstance(cell.value,str):cell.data_type='s'
        for cell in sheet[1]:
            cell.data_type='s';cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24354B')
            cell.alignment=Alignment(wrap_text=True)
        sheet.freeze_panes='A2'
        if columns:sheet.auto_filter.ref=sheet.dimensions
        for index in range(1,len(columns)+1):
            width=max(len(str(sheet.cell(row,index).value or '')) for row in range(1,min(sheet.max_row,100)+1))
            sheet.column_dimensions[get_column_letter(index)].width=min(48,max(14,width+2))
    if not book.worksheets:book.create_sheet('No records')
    result=io.BytesIO();book.save(result);return result.getvalue()
